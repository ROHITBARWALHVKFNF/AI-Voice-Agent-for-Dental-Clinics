
from io import BytesIO
from typing import Literal, Optional
from uuid import UUID

from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel, ConfigDict, Field
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.dependencies import get_current_user
from app.db.database import get_db
from app.models.knowledge_source import KnowledgeSource
from app.models.user import User


router = APIRouter(
    prefix="/knowledge-base/sources",
    tags=["Knowledge Base"],
)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
}

SOURCE_TYPES = {
    "text",
    "faq",
    "clinic_details",
    "website",
    "document",
}

SOURCE_STATUSES = {
    "draft",
    "needs_review",
    "approved",
}


class KnowledgeSourceCreate(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    category: str = Field(
        default="General information",
        min_length=1,
        max_length=100,
    )
    source_type: Literal[
        "text",
        "faq",
        "clinic_details",
        "website",
    ] = "text"
    content: str = Field(default="", max_length=500_000)
    source_url: Optional[str] = Field(default=None, max_length=2048)
    status: Literal["draft", "needs_review"] = "needs_review"


class KnowledgeSourceUpdate(BaseModel):
    title: Optional[str] = Field(default=None, min_length=1, max_length=200)
    category: Optional[str] = Field(default=None, min_length=1, max_length=100)
    content: Optional[str] = Field(default=None, max_length=500_000)
    source_url: Optional[str] = Field(default=None, max_length=2048)
    status: Optional[Literal["draft", "needs_review", "approved"]] = None


class KnowledgeSourceResponse(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: UUID
    organization_id: UUID
    title: str
    category: str
    source_type: str
    content: str
    source_url: Optional[str]
    original_filename: Optional[str]
    mime_type: Optional[str]
    size_bytes: Optional[int]
    status: str
    created_at: object
    updated_at: object


def get_source_or_404(
    db: Session,
    source_id: UUID,
    organization_id: UUID,
) -> KnowledgeSource:
    source = db.scalar(
        select(KnowledgeSource).where(
            KnowledgeSource.id == source_id,
            KnowledgeSource.organization_id == organization_id,
        )
    )

    if source is None:
        raise HTTPException(
            status_code=404,
            detail="Knowledge source not found.",
        )

    return source


def extract_document_text(filename: str, data: bytes) -> str:
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        try:
            reader = PdfReader(BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail="Could not read this PDF. Check that it is valid and not encrypted.",
            ) from exc

    elif lower_name.endswith(".docx"):
        try:
            document = Document(BytesIO(data))
            paragraphs = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(paragraphs)
        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail="Could not read this DOCX document.",
            ) from exc

    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(
                status_code=422,
                detail="This file must be encoded as UTF-8 text.",
            ) from exc

    text = text.strip()

    if not text:
        raise HTTPException(
            status_code=422,
            detail=(
                "No text could be extracted. Scanned PDFs need OCR, "
                "which is not implemented yet."
            ),
        )

    if len(text) > 500_000:
        raise HTTPException(
            status_code=413,
            detail="Extracted text exceeds the 500,000-character limit.",
        )

    return text


@router.get("", response_model=list[KnowledgeSourceResponse])
def list_sources(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return db.scalars(
        select(KnowledgeSource)
        .where(
            KnowledgeSource.organization_id == current_user.organization_id
        )
        .order_by(KnowledgeSource.updated_at.desc())
    ).all()


@router.post("", response_model=KnowledgeSourceResponse, status_code=201)
def create_source(
    payload: KnowledgeSourceCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    title = payload.title.strip()
    content = payload.content.strip()
    source_url = payload.source_url.strip() if payload.source_url else None

    if not title:
        raise HTTPException(status_code=422, detail="Title is required.")

    if payload.source_type == "website":
        if not source_url or not source_url.startswith(("https://", "http://")):
            raise HTTPException(
                status_code=422,
                detail="Provide a valid HTTP or HTTPS website URL.",
            )
        if content:
            raise HTTPException(
                status_code=422,
                detail="Website content must be extracted by a separate ingestion step.",
            )
    elif not content:
        raise HTTPException(
            status_code=422,
            detail="Content is required for this source type.",
        )

    source = KnowledgeSource(
        organization_id=current_user.organization_id,
        created_by_user_id=current_user.id,
        title=title,
        category=payload.category.strip(),
        source_type=payload.source_type,
        content=content,
        source_url=source_url,
        status=payload.status,
    )

    db.add(source)
    db.commit()
    db.refresh(source)
    return source


@router.post("/upload", response_model=list[KnowledgeSourceResponse], status_code=201)
def upload_documents(
    files: list[UploadFile] = File(...),
    category: str = "Documents",
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if not files:
        raise HTTPException(status_code=422, detail="Select at least one file.")

    if len(files) > 10:
        raise HTTPException(status_code=413, detail="Upload up to 10 files at a time.")

    created_sources = []

    try:
        for uploaded_file in files:
            filename = (uploaded_file.filename or "").strip()
            if not filename:
                raise HTTPException(status_code=422, detail="A file has no filename.")

            extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
            if extension not in ALLOWED_EXTENSIONS:
                raise HTTPException(
                    status_code=415,
                    detail=f"Unsupported file type for {filename}. Supported: PDF, DOCX, TXT, MD, CSV.",
                )

            data = uploaded_file.file.read(MAX_UPLOAD_BYTES + 1)

            if not data:
                raise HTTPException(
                    status_code=422,
                    detail=f"{filename} is empty.",
                )

            if len(data) > MAX_UPLOAD_BYTES:
                raise HTTPException(
                    status_code=413,
                    detail=f"{filename} exceeds the 10 MB limit.",
                )

            extracted_text = extract_document_text(filename, data)

            source = KnowledgeSource(
                organization_id=current_user.organization_id,
                created_by_user_id=current_user.id,
                title=filename[:200],
                category=category.strip()[:100] or "Documents",
                source_type="document",
                content=extracted_text,
                original_filename=filename[:255],
                mime_type=uploaded_file.content_type,
                size_bytes=len(data),
                status="needs_review",
            )
            db.add(source)
            created_sources.append(source)

        db.commit()

        for source in created_sources:
            db.refresh(source)

        return created_sources

    except Exception:
        db.rollback()
        raise

    finally:
        for uploaded_file in files:
            uploaded_file.file.close()


@router.patch("/{source_id}", response_model=KnowledgeSourceResponse)
def update_source(
    source_id: UUID,
    payload: KnowledgeSourceUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    source = get_source_or_404(
        db,
        source_id,
        current_user.organization_id,
    )

    changes = payload.model_dump(exclude_unset=True)

    for field, value in changes.items():
        if isinstance(value, str):
            value = value.strip()

        if field == "title" and not value:
            raise HTTPException(status_code=422, detail="Title cannot be empty.")

        if field == "content" and not value:
            raise HTTPException(status_code=422, detail="Content cannot be empty.")

        if field == "source_url" and value and not value.startswith(("https://", "http://")):
            raise HTTPException(
                status_code=422,
                detail="Website URLs must start with HTTP or HTTPS.",
            )

        setattr(source, field, value)

    db.commit()
    db.refresh(source)
    return source


@router.delete("/{source_id}")
def delete_source(
    source_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    source = get_source_or_404(
        db,
        source_id,
        current_user.organization_id,
    )

    db.delete(source)
    db.commit()

    return {"message": "Knowledge source deleted."}
